import io,os,random,secrets
from datetime import date
from fasthtml.common import *
from fastlite import *
from monsterui.all import *
import numpy as np
from openpyxl.styles import Font,PatternFill
import pandas as pd

def set_scope_before(req, sess):
    req.scope.setdefault('auth', sess.get('auth'))

def user_auth_before(req, sess):
    auth = req.scope.get('auth')
    if auth: return
    sess['next'] = req.url.path
    return Redirect('/login')
    
def mk_nav(auth=None):
    authed = [A(href='/tests/new')(H4('Test')), A(href='/tests')(H4('Data')), A(href='/logout')(H4('Logout'))]
    anon = [A(href='/login')(H4('Login'))]
    return NavBar(
        *(authed if auth else anon),
        brand=A(href='/')(H2('COBALT')),
        cls='bg-primary text-primary-content py-4 px-16')

def mk_shell(c, req):
    auth = req.scope.get('auth', None)
    return mk_nav(auth), Main(id='main')(c)

before = [
    Beforeware(set_scope_before),
    Beforeware(user_auth_before, skip=[r'/favicon\.ico', '/login', '/logout', '/'])]
    
app,rt = fast_app(
    title='COBALT',
    before=before,
    session_cookie='session_cobalt',
    secret_key=os.environ.get('SECRET_KEY', secrets.token_hex(32)),
    hdrs=Theme.blue.headers(katex=True, apex_charts=True),
    body_wrap=mk_shell)

@rt
def index(): return Container(render_md(Path('README.md').read_text()))

RESEARCHER_USER,RESEARCHER_PASS = os.getenv('RESEARCHER_USER'),os.getenv('RESEARCHER_PASS')

@rt('/login')
def get(sess):
    if sess.get('auth'): return Redirect('/')
    return Div(cls='flex justify-center pt-4')(
        Card(
            H3('Researcher Login'),
            P(id='errormsg', cls=TextT.error+TextT.bold),
            Form(hx_post='/login', hx_target='#errormsg')(
                LabelInput('Username', id='username', name='username', autocomplete='username'),
                LabelInput('Password', id='password', name='password', type='password', autocomplete='current-password'),
                Button(cls=ButtonT.primary)('Log in')),
            cls='w-96'))

@rt('/login')
def post(sess, username:str, password:str):
    if username==RESEARCHER_USER and password==RESEARCHER_PASS:
        sess['auth'] = True
        return Redirect(sess.pop('next', '/'))
    else: return 'Wrong username/password.'

@rt
def logout(sess):
    sess.clear()
    return Redirect('/')

db = database('data/cobalt.db')

class Test: id:int; initials:str; birthday:str; sex:str; test_date:str; threshold:float; completed:bool=False; srt_db:float=0.0; srt_done:bool=False
class Trial: id:int; test_id:int; trial_num:int; gap_ms:float; answer:str; correct:str
class SrtTrial: id:int; test_id:int; trial_num:int; level_db:float; word:str; correct:bool
tests_db = db.create(Test, pk='id')
trials_db = db.create(Trial, pk='id', foreign_keys=[('test_id', 'test', 'id')])
srt_trials_db = db.create(SrtTrial, pk='id', foreign_keys=[('test_id', 'test', 'id')])

def age_at(t):
    bday = date.fromisoformat(t.birthday)
    tday = date.fromisoformat(t.test_date)
    return tday.year-bday.year-((tday.month,tday.day)<(bday.month,bday.day))

def test_sem(tid):
    revs = replay(get_trials(tid))[3][-4:]
    vals = np.array([gap for _,gap in revs])
    return vals.std(ddof=1)/np.sqrt(len(vals)) if len(vals)>1 else 0

def sort_th(col, sort, dir):
    new_dir = 'desc' if col==sort and dir=='asc' else 'asc'
    mark = (' ↑' if dir=='asc' else ' ↓') if col==sort else ''
    return Th(A(hx_get=f'/tests?sort={col}&dir={new_dir}', hx_target='#main', hx_push_url='true')(col+mark))

def row_link(tid, v): return Td(A(href=f'/tests/{tid}', cls='block w-full')(v))

def sorted_tests(sort, dir):
    tsts = list(tests_db())
    rev = dir == 'desc'
    if sort == 'age': return sorted(tsts, key=age_at, reverse=rev)
    if sort == 'sem':
        done = [t for t in tsts if t.completed]
        pending = [t for t in tsts if not t.completed]
        return sorted(done, key=lambda t:test_sem(t.id), reverse=rev)+pending
    return list(tests_db(order_by=f'{sort} {"desc" if rev else "asc"}'))

@rt('/tests')
def get(sort:str='id', dir:str='desc'):
    if not tests_db(): return Div(cls='flex justify-center pt-8')(
        Card(
            H3('No tests found'),
            P('Create ', A(href='/tests/new', cls=AT.primary)('a new test'), ' to get started.'),
            cls=TextT.center))
    cols = [*tests_db.columns_dict.keys(), 'age', 'sem']
    if sort not in cols: sort = 'id'
    tsts = sorted_tests(sort, dir)
    def val(t, col):
        if col=='age': return f'{age_at(t)} y'
        if col=='sem': return f'{test_sem(t.id):.1f} ms' if t.completed else 'pending'
        if col=='threshold': return f'{t.threshold:.1f}' if t.completed else 'pending'
        return getattr(t, col)
    tests_table = Table(
        Thead(Tr(*[sort_th(col, sort, dir) for col in cols])),
        Tbody(*[Tr(*[row_link(t.id, val(t, col)) for col in cols]) for t in tsts]))
    return Div(cls='flex flex-col gap-4 p-4')(
        DivRAligned(A(href='/tests/export', cls=ButtonT.primary)('Export to Excel')),
        tests_table)


STEP,START_GAP,N_REVERSALS = 0.8,100.0,6

def step_state(gap, streak, last_dir, ok):
    "Advance staircase one trial; returns new gap, streak, direction, and whether it reversed"
    new_dir = None
    if ok:
        streak += 1
        if streak>1: gap,streak,new_dir = gap*STEP,0,'down'
    else: gap,streak,new_dir = gap/STEP,0,'up'
    return gap,streak,(new_dir or last_dir),bool(new_dir and last_dir and new_dir!=last_dir)

def replay(trls):
    "Derive current staircase state from a test's trials"
    gap,streak,last_dir,revs = START_GAP,0,None,[]
    for t in trls:
        prev = gap
        gap,streak,last_dir,rev = step_state(gap, streak, last_dir, t.answer==t.correct)
        if rev: revs.append((t.trial_num, prev))
    return gap,streak,last_dir,revs

def get_trials(tid): return trials_db(where='test_id=?', where_args=[tid], order_by='trial_num')

@rt('/tests/new')
def get():
    return Div(cls='pt-4 flex justify-center')(
        Card(
            H3('New test'),
            Form(method='post', action='/tests')(
                LabelInput('Initials', id='initials', name='initials', required=True),
                LabelInput('Birthday', id='birthday', name='birthday', type='date', required=True),
                Fieldset(
                    Legend(cls=TextT.medium)('Sex'),
                    DivHStacked(
                        LabelRadio('Male', id='sex-m', name='sex', value='M'),
                        LabelRadio('Female', id='sex-f', name='sex', value='F'),
                        cls='gap-4 p-4')),
                Button(cls=ButtonT.primary)('Start test')),
            cls='w-96'))

@rt('/tests')
def post(initials:str, birthday:str, sex:str):
    t = tests_db.insert(Test(initials=initials, birthday=birthday, sex=sex, test_date=str(date.today()), threshold=0.0))
    return Redirect(f'/tests/{t.id}')

DEFAULT_WIDTH = 16
COLUMN_WIDTHS = dict(README=dict(field=32, description=48))
COBALT = PatternFill('solid', fgColor='2F5597')
OFF_WHITE = PatternFill('solid', fgColor='F8FAFC')
HEADER_FONT = Font(color='FFFFFF', bold=True)
DATA_HEADERS = dict(
    tests=dict(
        test_id='Test ID', initials='Initials', birthday='Birthday', sex='Sex', test_date='Test date', age='Age at test', threshold_ms='Threshold (ms)', sem_ms='SEM (ms)', completed='Completed', srt_db='SRT level (dB)', srt_done='SRT completed'),
    trials=dict(
        test_id='Test ID', initials='Initials', test_date='Test date', age='Age at test', sex='Sex', trial_num='Trial number', gap_ms='Gap (ms)', answer='Answer', correct='Correct answer', is_correct='Correct', reversal='Reversal'),
    srt_trials=dict(
        test_id='Test ID', initials='Initials', test_date='Test date', age='Age at test', sex='Sex', trial_num='Trial number', level_db='Level (dB)', word='Word', is_correct='Correct'))
README_FIELDS = dict(
    exported='Export date',
    START_GAP='Initial TOJ gap',
    STEP='Staircase multiplier',
    N_REVERSALS='Reversals until test ends',
    threshold_ms='Threshold',
    sem_ms='Standard error of the mean',
    age='Age',
    reversal='Reversal',
    is_correct='Correct response',
    blank_cells='Unavailable measure')

def test_row(t):
    return dict(
        test_id=t.id, initials=t.initials, birthday=t.birthday, sex=t.sex, test_date=t.test_date, age=age_at(t),
        threshold_ms=t.threshold if t.completed else None, sem_ms=test_sem(t.id) if t.completed else None,
        completed=int(t.completed), srt_db=t.srt_db, srt_done=int(t.srt_done))

def test_keys(r): return {k:r[k] for k in ('test_id','initials','test_date','age','sex')}

def trial_rows(t, keys):
    trl = get_trials(t.id)
    revn = {n for n,_ in replay(trl)[3]}
    return [keys | dict(trial_num=o.trial_num, gap_ms=o.gap_ms, answer=o.answer, correct=o.correct, is_correct=int(o.answer==o.correct), reversal=int(o.trial_num in revn)) for o in trl]

def srt_rows(t, keys):
    trl = srt_trials_db(where='test_id=?', where_args=[t.id], order_by='trial_num')
    return [keys | dict(trial_num=o.trial_num, level_db=o.level_db, word=o.word, is_correct=int(o.correct)) for o in trl]

def readme_rows():
    return [dict(field=README_FIELDS['exported'], description=str(date.today())),
            dict(field=README_FIELDS['START_GAP'], description=f'{START_GAP} ms'),
            dict(field=README_FIELDS['STEP'], description=STEP),
            dict(field=README_FIELDS['N_REVERSALS'], description=N_REVERSALS),
            dict(field=README_FIELDS['threshold_ms'], description='Mean gap of the last 4 reversals'),
            dict(field=README_FIELDS['sem_ms'], description='SEM of the last 4 reversal gaps'),
            dict(field=README_FIELDS['age'], description='Years from birthday to test date'),
            dict(field=README_FIELDS['reversal'], description='1 if the trial ended a staircase run'),
            dict(field=README_FIELDS['is_correct'], description='1 if the response matched the stimulus'),
            dict(field=README_FIELDS['blank_cells'], description='Measure not yet available')]

def set_widths(ws, widths):
    for col in ws.iter_cols():
        name = col[0].value
        ws.column_dimensions[col[0].column_letter].width = widths.get(name, DEFAULT_WIDTH)

def style_sheet(ws, row_header=False):
    for cell in ws[1]:
        cell.fill = COBALT
        cell.font = HEADER_FONT
    for row in ws.iter_rows(min_row=2):
        for cell in row: cell.fill = OFF_WHITE
    if row_header:
        for row in ws.iter_rows(min_row=2, max_col=1):
            row[0].fill = COBALT
            row[0].font = HEADER_FONT

def export_book():
    tsts = list(tests_db())
    rows = [test_row(t) for t in tsts]
    trls = [r for t,r in zip(tsts, rows) for r in trial_rows(t, test_keys(r))]
    srts = [r for t,r in zip(tsts, rows) for r in srt_rows(t, test_keys(r))]
    sheets = dict(tests=rows, trials=trls, srt_trials=srts)
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine='openpyxl') as xl:
        pd.DataFrame(readme_rows()).to_excel(xl, sheet_name='README', index=False)
        ws = xl.sheets['README']
        style_sheet(ws, row_header=True)
        set_widths(ws, COLUMN_WIDTHS['README'])
        for name,data in sheets.items():
            df = pd.DataFrame(data).rename(columns=DATA_HEADERS[name])
            df.to_excel(xl, sheet_name=name, index=False)
            ws = xl.sheets[name]
            style_sheet(ws)
            set_widths(ws, COLUMN_WIDTHS.get(name, {}))
            ws.freeze_panes = 'A2'
    return buf.getvalue()

@rt('/tests/export')
def get():
    fn = f'cobalt_{date.today()}.xlsx'
    return Response(export_book(), media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'content-disposition': f'attachment; filename="{fn}"'})

def plot(x, y, xlabel='', ylabel='', title='', precision=0, cls=None, id=None, name=None, annotations=None, markers=None, **kwargs):
    chart = dict(type='line', zoom=dict(enabled=False), toolbar=dict(show=False), height=380) | kwargs
    opts = dict(
        chart=chart,
        series=[dict(data=y, **({'name':name} if name else {}))],
        xaxis=dict(title=dict(text=xlabel), categories=x),
        yaxis=dict(title=dict(text=ylabel), decimalsInFloat=precision))
    if title: opts['title'] = dict(text=title)
    if annotations: opts['annotations'] = annotations
    if markers: opts['markers'] = markers
    return ApexChart(opts=opts, cls=cls, id=id)

def patient_cards(t, **attrs):
    return DivHStacked(
        Card(
            P(f'Initials: {t.initials}'),
            P(f'Birthday: {t.birthday}'),
            P(f'Age at test: {age_at(t)}'),
            P(f'Sex: {t.sex}'),
            cls='w-full'),
        Card(
            P(f'Threshold: {t.threshold:.1f} ± {test_sem(t.id):.1f} ms' if t.completed else 'Threshold: pending'),
            P(f'Test date: {t.test_date}'),
            P(f'Status: {"Completed" if t.completed else "In progress"}'),
            cls='w-full'),
        cls='items-stretch gap-4', id='patient-card', **attrs)

def gap_chart(trl, revs, next_num=None, next_gap=None):
    x,y = [t.trial_num for t in trl],[t.gap_ms for t in trl]
    if next_num is not None: x.append(next_num); y.append(next_gap)
    used = revs[-4:]
    if len(used)<4: return plot(x, y, 'Trial number', 'Gap (ms)', precision=0, name='Gap (ms)', cls='w-full')
    vals = np.array([g for _,g in used])
    th,sem = vals.mean(),vals.std(ddof=1)/np.sqrt(len(vals))
    idxs = [x.index(n) for n,_ in used]
    markers = dict(size=0, discrete=[dict(seriesIndex=0, dataPointIndex=i, size=7, fillColor='#dc2626', strokeColor='#fff', strokeWidth=2) for i in idxs])
    annotations = dict(
        yaxis=[
            dict(y=th-sem, y2=th+sem, fillColor='#ef4444', opacity=.15, label=dict(text=f'{th:.1f} ± {sem:.1f} ms', style=dict(background='#ef4444', color='#fff'))),
            dict(y=th, borderColor='#ef4444', strokeDashArray=0)])
    return plot(x, y, 'Trial number', 'Gap (ms)', precision=0, name='Gap (ms)', cls='w-full', annotations=annotations, markers=markers)

def done_card(t, n):
    return Card(
        H4('Test complete'),
        P(f'Trials: {n}'),
        P(cls=TextPresets.bold_lg)(f'Threshold: {t.threshold:.1f} ± {test_sem(t.id):.1f} ms'),
        A(href='/tests/new', cls=AT.primary)('New test'))

def css_arrow(dir, w='[2rem]', h='[0.5rem]'):
    side = 'l' if dir=='r' else 'r'
    return Span(cls=f'border-y-{h} border-y-transparent border-{side}-{w} border-{side}-current')

def trial_card(tid, num, gap, revs, side, lvl):
    btn_attrs = dict(type='button', name='answer', hx_post=f'/tests/{tid}/trials', hx_target='#test-panel', hx_vals=f'{{"correct":"{side}"}}')
    btn_cls = 'flex-auto py-8 gap-2'
    return Card(
        H4(f'Trial {num} — gap {gap:.1f} ms — reversals {len(revs)}/{N_REVERSALS}'),
        P('Put on headphones. Click play, then choose the ear that heard the sound first.'),
        Button(type='button', autofocus=True, onclick=f'playTOJ("{side}", {gap}, {lvl})')('Play sound'),
        DivHStacked(
            Button(**btn_attrs, value='left', cls=ButtonT.destructive+btn_cls, hx_trigger="click, keyup[key=='ArrowLeft'] from:body")(
                css_arrow('l'), P(cls=TextT.xl)('Left')),
            Button(**btn_attrs, value='right', cls=ButtonT.primary+btn_cls, hx_trigger="click, keyup[key=='ArrowRight'] from:body")(
                P(cls=TextT.xl)('Right'), css_arrow('r')),
            cls='gap-4 pt-4 w-full'))

def test_panel(tid):
    t = tests_db[tid]
    trl = get_trials(tid)
    gap,_,_,revs = replay(trl)
    panel = done_card(t, len(trl)) if t.completed else trial_card(tid, len(trl)+1, gap, revs, random.choice(['left','right']), toj_gain(t.srt_db))
    chart = gap_chart(trl, revs) if t.completed else gap_chart(trl, revs, len(trl)+1, gap)
    panel = done_card(t, len(trl)) if t.completed else trial_card(tid, len(trl)+1, gap, revs, random.choice(['left','right']), toj_gain(t.srt_db))
    return Div(id='test-panel', cls='flex flex-col gap-4')(panel, chart)

@rt('/tests/{tid}')
def get(tid:int):
    try: t = tests_db[tid]
    except NotFoundError: return Response(status_code=404)
    hdr = (Script(src='/js/playStimulus.js'), patient_cards(t))
    if not t.srt_done: return Container(*hdr, srt_panel(tid), cls='flex flex-col gap-4 p-8')
    if not srt_ok(t.srt_db): return Container(*hdr, headroom_card(tid, t), cls='flex flex-col gap-4 p-8')
    return Container(*hdr, test_panel(tid), cls='flex flex-col gap-4 p-8')

@rt('/tests/{tid}/trials')
def post(tid:int, answer:str, correct:str):
    t = tests_db[tid]
    if t.completed: return test_panel(tid)
    trl = get_trials(tid)
    gap,streak,last_dir,revs = replay(trl)
    ngap,_,_,rev = step_state(gap, streak, last_dir, answer==correct)
    trials_db.insert(Trial(test_id=tid, trial_num=len(trl)+1, gap_ms=gap, answer=answer, correct=correct))
    if rev: revs.append((len(trl)+1, gap))
    if len(revs)>=N_REVERSALS:
        tests_db.update(dict(id=tid, threshold=sum(g for _,g in revs[-4:])/4, completed=True))
        return test_panel(tid), patient_cards(tests_db[tid], hx_swap_oob='true')
    return test_panel(tid)

def srt_panel(tid):
    t = tests_db[tid]
    trl = get_srt_trials(tid)
    lvl,_,revs = srt_replay(trl)
    return Div(id='srt-panel', cls='flex flex-col gap-4')(
        srt_card(tid, len(trl)+1, lvl, next_word(trl), revs),
        srt_chart(trl))

def srt_chart(trl): return plot([t.trial_num for t in trl], [t.level_db for t in trl], 'Trial number', 'Level (dB)', precision=0, name='Level (dB)', cls='w-full')

def srt_card(tid, num, lvl, word, revs):
    return Card(
        H4(f'Calibration trial {num} — {lvl:.0f} dB — reversals {len(revs)}/{SRT_REVERSALS}'),
        P('Play the word, then mark whether the patient repeated it correctly.'),
        P(cls='text-3xl font-bold tracking-wide py-2')(word),
        Button(type='button', onclick=f'playSpondee("{word}", {lvl})', cls=ButtonT.secondary)('Play word'),
        Form(hx_post=f'/tests/{tid}/srt', hx_target='#srt-panel', hx_swap='outerHTML')(
            Input(type='hidden', name='word', value=word),
            Input(type='hidden', name='level_db', value=str(lvl)),
            DivHStacked(
                Button(name='correct', value='1', cls=ButtonT.primary)('Correct'),
                Button(name='correct', value='0', cls=ButtonT.secondary)('Incorrect'),
                cls='gap-4 pt-4')))

def headroom_card(tid, t):
    return Card(
        H4(cls=TextT.error+TextT.bold)('Insufficient headroom'),
        P(f'SRT came out at {t.srt_db:.1f} dB, so the TOJ level would be {toj_gain(t.srt_db):.1f} dB — above full scale.'),
        P('Set the tablet volume to maximum and recalibrate.'),
        Form(method='post', action=f'/tests/{tid}/recalibrate')(Button(cls=ButtonT.primary)('Recalibrate')))

@rt('/tests/{tid}/srt')
def post(tid:int, word:str, level_db:float, correct:int):
    t = tests_db[tid]
    if t.srt_done: return
    trl = get_srt_trials(tid)
    lvl,last_dir,revs = srt_replay(trl)
    nlvl,ndir,rev = srt_step(lvl, last_dir, len(revs), bool(correct))
    srt_trials_db.insert(SrtTrial(test_id=tid, trial_num=len(trl)+1, level_db=level_db, word=word, correct=bool(correct)))
    if rev: revs.append(nlvl)
    if len(revs)>=SRT_REVERSALS:
        tests_db.update(dict(id=tid, srt_db=sum(revs[-4:])/4, srt_done=True))
        t = tests_db[tid]
        panel = headroom_card(tid, t) if not srt_ok(t.srt_db) else test_panel(tid)
        return Div(id='srt-panel')(panel), patient_cards(t, hx_swap_oob='true')
    return srt_panel(tid)

@rt('/tests/{tid}/recalibrate')
def post(tid:int):
    srt_trials_db.delete_where('test_id=?', [tid])
    tests_db.update(dict(id=tid, srt_db=0.0, srt_done=False))
    return Redirect(f'/tests/{tid}')

SRT_START,SRT_BIG,SRT_SMALL,SRT_REVERSALS,SRT_HEADROOM = -30.0,5.0,2.0,6,50.0
SPONDEES = ('abelha boneca cadeira').split()

def srt_step(lvl, last_dir, nrev, ok):
    "Advance SRT staircase one trial; returns new level, direction, and whether it reversed"
    step = SRT_BIG if nrev<1 else SRT_SMALL
    new_dir = 'down' if ok else 'up'
    return (lvl-step if ok else lvl+step),new_dir,bool(last_dir and new_dir!=last_dir)

def srt_replay(trl):
    "Derive current SRT staircase state from a test's trials"
    lvl,last_dir,revs = SRT_START,None,[]
    for t in trl:
        lvl,last_dir,rev = srt_step(lvl, last_dir, len(revs), t.correct)
        if rev: revs.append(lvl)
    return lvl,last_dir,revs

def get_srt_trials(tid): return srt_trials_db(where='test_id=?', where_args=[tid], order_by='trial_num')
def next_word(trl): return random.choice([w for w in SPONDEES if w not in {t.word for t in trl}] or SPONDEES)
def toj_gain(srt_db): return srt_db + SRT_HEADROOM
def srt_ok(srt_db): return srt_db <= -SRT_HEADROOM

serve()
